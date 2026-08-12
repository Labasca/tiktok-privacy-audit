#!/usr/bin/env python3
"""
Turn the remediation board into a spreadsheet.

The board lives in the published artifact and its cards are the source of truth,
so this reads the JSON extracted from that page rather than a second copy of the
list. Keeping one source is the whole point: a hand-maintained sheet drifts from
the board within a week.

Output: tiktok-remediation-board.xlsx
  - one row per card, one step per line inside the Steps cell
  - freeze panes under the header, autofilter on every column
  - Status and Tag colour-coded, priority colour-coded
  - column widths and wrapping set so it is readable the moment it opens

Usage:  .venv\\Scripts\\python.exe build_board_xlsx.py [cards.json]
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "board-cards.json")
OUT = os.path.join(HERE, "tiktok-remediation-board.xlsx")

COL_NAME = {"fix": "Fix now", "gray": "Gray area", "good": "Already good"}
TAG_NAME = {"exp": "Exposed", "link": "Linked", "hard": "Hard", "clean": "Clean"}

# The artifact's palette, so the sheet and the board read as one thing.
INK = "0D1416"
STATUS_FILL = {"Fix now": "F6E6D2", "Gray area": "DDE7F2", "Already good": "D9EEE3"}
STATUS_FONT = {"Fix now": "8A5807", "Gray area": "215E9C", "Already good": "0F6749"}
TAG_FILL = {"Exposed": "F6E6D2", "Linked": "E5DEF7", "Hard": "F7DFDC", "Clean": "D9EEE3"}
TAG_FONT = {"Exposed": "8A5807", "Linked": "5C43B4", "Hard": "A8322A", "Clean": "0F6749"}
PRI_FONT = {"P1": "A8322A", "P2": "8A5807", "P3": "77898B"}

HEAD_FILL = PatternFill("solid", fgColor="E3E9E8")
THIN = Side(style="thin", color="D3DCDB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADERS = [
    ("Status", 14),
    ("Priority", 9),
    ("Tag", 11),
    ("Check", 44),
    ("What the audit found", 74),
    ("Steps", 60),
    ("Owner", 14),
    ("Notes", 30),
]


def build(cards):
    wb = Workbook()
    ws = wb.active
    ws.title = "Board"

    for i, (name, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = Font(bold=True, size=10, color=INK, name="Calibri")
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 24

    # Fix now first, then by priority, then by the order they sit on the board.
    order = {"fix": 0, "gray": 1, "good": 2}
    rows = sorted(enumerate(cards), key=lambda p: (order[p[1]["col"]], p[1]["pri"], p[0]))

    for r, (_, c) in enumerate(rows, start=2):
        status = COL_NAME[c["col"]]
        tag = TAG_NAME.get((c["tags"] or ["exp"])[0], "Exposed")
        pri = "P%d" % c["pri"]
        steps = "\n".join("%d. %s" % (i, s) for i, s in enumerate(c["steps"], 1))

        values = [status, pri, tag, c["title"], c["note"], steps, "", ""]
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=i in (4, 5, 6, 8),
                horizontal="center" if i in (2,) else "left",
            )
            cell.font = Font(size=10, name="Calibri", color=INK)

        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=STATUS_FILL[status])
        ws.cell(row=r, column=1).font = Font(size=10, bold=True, name="Calibri",
                                             color=STATUS_FONT[status])
        ws.cell(row=r, column=2).font = Font(size=10, bold=True, name="Calibri",
                                             color=PRI_FONT[pri])
        ws.cell(row=r, column=3).fill = PatternFill("solid", fgColor=TAG_FILL[tag])
        ws.cell(row=r, column=3).font = Font(size=10, bold=True, name="Calibri",
                                             color=TAG_FONT[tag])
        ws.cell(row=r, column=4).font = Font(size=10, bold=True, name="Calibri", color=INK)
        ws.cell(row=r, column=5).font = Font(size=9.5, name="Calibri", color="495A5C")
        ws.cell(row=r, column=6).font = Font(size=9.5, name="Calibri", color="495A5C")

        # Roughly one line per 95 characters of the longest wrapped column, so a
        # long note does not open as a single clipped line.
        lines = max(len(c["note"]) // 95 + 1, len(c["steps"]) or 1, len(c["title"]) // 40 + 1)
        ws.row_dimensions[r].height = min(15 * lines + 6, 190)

    last = len(rows) + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:H%d" % last

    # Dropdowns, so editing in the sheet cannot invent a status or a tag that the
    # board has no column or colour for.
    def dv(formula, col):
        v = DataValidation(type="list", formula1=formula, allow_blank=False)
        ws.add_data_validation(v)
        v.add("%s2:%s%d" % (col, col, last))

    dv('"Fix now,Gray area,Already good"', "A")
    dv('"P1,P2,P3"', "B")
    dv('"Exposed,Linked,Hard,Clean"', "C")

    # A second sheet holding what the tags mean, because the sheet will outlive
    # the conversation that explains them.
    key = wb.create_sheet("Key")
    key.column_dimensions["A"].width = 16
    key.column_dimensions["B"].width = 96
    rows_key = [
        ("Column", "What it means"),
        ("Fix now", "Ours to change, and worth changing. Sorted by priority."),
        ("Gray area", "Either no clean fix, or not measured yet. Parked, not ignored."),
        ("Already good", "We look like an ordinary phone here, or the check tells them nothing."),
        ("", ""),
        ("Tag", "What it means"),
        ("Exposed", "We look wrong here today and it is ours to fix."),
        ("Linked", "Fine on one phone, ties the whole fleet together at scale."),
        ("Hard", "True on our rack and no software fix exists on stock iOS."),
        ("Clean", "We already look normal, or the read carries no signal about us."),
        ("", ""),
        ("Priority", "What it means"),
        ("P1", "Do this first. Cheap, high signal, or both."),
        ("P2", "Real work, real payoff, not urgent this week."),
        ("P3", "Worth doing once the P1s and P2s are clear."),
        ("", ""),
        ("Source", "Frida captures 4 to 6 August 2026, 15 runs, 3 of them published posts."),
        ("", "iPhone X on iOS 16.7.16, TikTok 46.3.0. Read-only instrumentation."),
        ("", "Tags are our judgement about our own rack, not a verdict from the app."),
    ]
    for r, (a, b) in enumerate(rows_key, start=1):
        ca, cb = key.cell(row=r, column=1, value=a), key.cell(row=r, column=2, value=b)
        head = b == "What it means" or a == "Source"
        ca.font = Font(bold=True, size=10, name="Calibri",
                       color=INK if head else TAG_FONT.get(a, INK))
        cb.font = Font(bold=head, size=10, name="Calibri", color=INK if head else "495A5C")
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        if head:
            ca.fill = cb.fill = HEAD_FILL

    wb.save(OUT)
    return len(rows)


if __name__ == "__main__":
    with open(SRC, encoding="utf-8") as f:
        cards = json.load(f)
    n = build(cards)
    print("wrote %s  (%d rows)" % (OUT, n))
